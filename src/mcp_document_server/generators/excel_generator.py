"""Excel spreadsheet generation using openpyxl."""

import logging
import re
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

logger = logging.getLogger(__name__)


class ExcelGenerator:
    """Generate Excel spreadsheets from structured data."""

    def __init__(self) -> None:
        """Initialize the Excel generator."""
        self.logger = logging.getLogger(f"{__name__}.{self.__class__.__name__}")

    def generate_spreadsheet(
        self,
        sheets: list[dict[str, Any]],
        output_path: Path,
        include_formulas: bool = True,
    ) -> Path:
        """
        Generate an Excel spreadsheet from sheet data.

        Args:
            sheets: List of worksheet dictionaries with name and data
            output_path: Path where the spreadsheet should be saved
            include_formulas: Whether to process formula strings

        Returns:
            Path to the generated spreadsheet

        Raises:
            ValueError: If sheet data is invalid
        """
        try:
            self.logger.info(
                f"Generating Excel spreadsheet with {len(sheets)} sheet(s), "
                f"formulas={'enabled' if include_formulas else 'disabled'}"
            )

            # Create workbook
            wb = Workbook()
            if wb.active:
                wb.remove(wb.active)  # Remove default sheet

            # Process each sheet
            for idx, sheet_data in enumerate(sheets):
                self._add_sheet(wb, sheet_data, idx, include_formulas)

            # Save workbook
            output_path.parent.mkdir(parents=True, exist_ok=True)
            wb.save(str(output_path))

            self.logger.info(f"Excel spreadsheet saved to {output_path}")
            return output_path

        except Exception as e:
            self.logger.error(f"Error generating Excel spreadsheet: {e}", exc_info=True)
            raise

    def _add_sheet(
        self, wb: Workbook, sheet_data: dict[str, Any], idx: int, include_formulas: bool
    ) -> None:
        """Add a worksheet to the workbook."""
        sheet_name = sheet_data.get("name", f"Sheet{idx + 1}")
        ws = wb.create_sheet(title=sheet_name)

        self.logger.debug(f"Adding sheet '{sheet_name}'")

        # Get sheet configuration
        headers = sheet_data.get("headers", [])
        data_rows = sheet_data.get("data", [])
        auto_size_columns = sheet_data.get("auto_size_columns", True)
        freeze_header = sheet_data.get("freeze_header", True)
        styling = sheet_data.get("styling", {})

        # Add headers
        if headers:
            self._add_headers(ws, headers, styling)

        # Add data rows
        start_row = 2 if headers else 1
        for row_idx, row_data in enumerate(data_rows):
            self._add_data_row(ws, row_data, start_row + row_idx, include_formulas)

        # Apply auto-sizing
        if auto_size_columns:
            self._auto_size_columns(ws)

        # Freeze header row
        if freeze_header and headers:
            ws.freeze_panes = "A2"

    def _add_headers(self, ws: Any, headers: list[str], styling: dict[str, Any]) -> None:
        """Add header row with styling."""
        for col_idx, header in enumerate(headers, start=1):
            cell = ws.cell(row=1, column=col_idx)
            cell.value = header

            # Apply header styling
            cell.font = Font(bold=True, size=12, color="FFFFFF")
            cell.fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = Border(
                bottom=Side(style="thick", color="000000"),
                left=Side(style="thin", color="000000"),
                right=Side(style="thin", color="000000"),
            )

    def _add_data_row(
        self, ws: Any, row_data: list[Any], row_idx: int, include_formulas: bool
    ) -> None:
        """Add a data row to the worksheet."""
        for col_idx, cell_data in enumerate(row_data, start=1):
            cell = ws.cell(row=row_idx, column=col_idx)

            # Check if cell contains a formula
            if include_formulas and isinstance(cell_data, str) and cell_data.startswith("="):
                cell.value = cell_data
            else:
                # Handle different data types
                if isinstance(cell_data, int | float):
                    cell.value = cell_data
                    cell.number_format = self._get_number_format(cell_data)
                elif isinstance(cell_data, str):
                    # Try to parse as number
                    try:
                        if "." in cell_data:
                            cell.value = float(cell_data)
                            cell.number_format = "#,##0.00"
                        else:
                            cell.value = int(cell_data)
                            cell.number_format = "#,##0"
                    except ValueError:
                        # Keep as string
                        cell.value = cell_data
                else:
                    cell.value = str(cell_data)

            # Apply basic styling
            cell.alignment = Alignment(vertical="center")
            cell.border = Border(
                left=Side(style="thin", color="D0D0D0"),
                right=Side(style="thin", color="D0D0D0"),
                bottom=Side(style="thin", color="D0D0D0"),
            )

            # Apply alternating row colors
            if row_idx % 2 == 0:
                cell.fill = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")

    def _get_number_format(self, value: Any) -> str:
        """Determine the appropriate number format."""
        if isinstance(value, float):
            return "#,##0.00"
        elif isinstance(value, int):
            return "#,##0"
        return "General"

    def _auto_size_columns(self, ws: Any) -> None:
        """Auto-size columns based on content."""
        for column in ws.columns:
            max_length = 0
            column_letter = get_column_letter(column[0].column)

            for cell in column:
                try:
                    if cell.value:
                        cell_length = len(str(cell.value))
                        if cell_length > max_length:
                            max_length = cell_length
                except Exception:
                    pass

            # Set column width (add some padding)
            adjusted_width = min(max_length + 2, 50)  # Cap at 50 characters
            ws.column_dimensions[column_letter].width = adjusted_width

    def _parse_cell_reference(self, formula: str) -> list[str]:
        """Parse cell references from a formula string."""
        # Match cell references like A1, B2, etc.
        pattern = r"[A-Z]+\d+"
        return re.findall(pattern, formula)
