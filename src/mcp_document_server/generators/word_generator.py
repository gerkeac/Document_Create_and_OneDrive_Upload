"""Word document generation using python-docx."""

from datetime import datetime
from pathlib import Path
from typing import Any

from docx import Document

from mcp_document_server.utils.markdown_parser import (
    parse_inline_formatting,
    parse_markdown,
)


class WordGenerator:
    """Generate Word documents from markdown or JSON content."""

    def __init__(self) -> None:
        """Initialize the Word document generator."""
        self.doc: Document | None = None

    def generate_from_markdown(
        self,
        content: str,
        output_path: Path,
        title: str | None = None,
        author: str | None = None,
    ) -> Path:
        """
        Generate a Word document from markdown content.

        Args:
            content: Markdown string
            output_path: Path where the document should be saved
            title: Optional document title
            author: Optional document author

        Returns:
            Path to the generated document
        """
        self.doc = Document()

        # Set document metadata
        if title:
            self.doc.core_properties.title = title
        if author:
            self.doc.core_properties.author = author
        self.doc.core_properties.created = datetime.now()

        # Parse markdown content
        blocks = parse_markdown(content)

        # Generate document content
        for block in blocks:
            self._add_block(block)

        # Save the document
        output_path.parent.mkdir(parents=True, exist_ok=True)
        self.doc.save(str(output_path))

        return output_path

    def generate_from_json(
        self,
        content: dict[str, Any],
        output_path: Path,
    ) -> Path:
        """
        Generate a Word document from JSON content.

        Expected JSON structure:
        {
            "title": "Document Title",
            "author": "Author Name",
            "blocks": [
                {"type": "heading", "level": 1, "text": "Heading"},
                {"type": "paragraph", "text": "Paragraph text"},
                {"type": "list", "list_type": "bullet", "items": ["Item 1", "Item 2"]},
                {"type": "table", "headers": ["Col1", "Col2"], "rows": [["A", "B"]]}
            ]
        }

        Args:
            content: Structured JSON content
            output_path: Path where the document should be saved

        Returns:
            Path to the generated document
        """
        self.doc = Document()

        # Set document metadata
        if "title" in content:
            self.doc.core_properties.title = content["title"]
        if "author" in content:
            self.doc.core_properties.author = content["author"]
        self.doc.core_properties.created = datetime.now()

        # Process blocks
        blocks = content.get("blocks", [])
        for block in blocks:
            self._add_block(block)

        # Save the document
        output_path.parent.mkdir(parents=True, exist_ok=True)
        self.doc.save(str(output_path))

        return output_path

    def _add_block(self, block: dict[str, Any]) -> None:
        """
        Add a content block to the document.

        Args:
            block: Block dictionary with type and properties
        """
        if not self.doc:
            raise RuntimeError("Document not initialized")

        block_type = block.get("type")

        if block_type == "heading":
            self._add_heading(block)
        elif block_type == "paragraph":
            self._add_paragraph(block)
        elif block_type == "list":
            self._add_list(block)
        elif block_type == "table":
            self._add_table(block)
        else:
            # Unknown block type, treat as paragraph
            self._add_paragraph({"text": str(block)})

    def _add_heading(self, block: dict[str, Any]) -> None:
        """Add a heading to the document."""
        if not self.doc:
            return

        level = block.get("level", 1)
        text = block.get("text", "")

        # python-docx heading levels: 0=Title, 1-9=Heading 1-9
        # Our markdown levels: 1-6 for H1-H6
        self.doc.add_heading(text, level=level)

    def _add_paragraph(self, block: dict[str, Any]) -> None:
        """Add a paragraph to the document with inline formatting."""
        if not self.doc:
            return

        text = block.get("text", "")
        paragraph = self.doc.add_paragraph()

        # Parse inline formatting
        runs = parse_inline_formatting(text)

        for run_data in runs:
            run = paragraph.add_run(run_data.get("text", ""))
            if run_data.get("bold"):
                run.bold = True
            if run_data.get("italic"):
                run.italic = True
            if run_data.get("underline"):
                run.underline = True

    def _add_list(self, block: dict[str, Any]) -> None:
        """Add a list to the document."""
        if not self.doc:
            return

        list_type = block.get("list_type", "bullet")
        items = block.get("items", [])

        for item in items:
            paragraph = self.doc.add_paragraph(style="List Bullet")

            # For numbered lists, use List Number style
            if list_type == "number":
                paragraph.style = "List Number"

            # Add the item text with inline formatting
            runs = parse_inline_formatting(item)
            for run_data in runs:
                run = paragraph.add_run(run_data.get("text", ""))
                if run_data.get("bold"):
                    run.bold = True
                if run_data.get("italic"):
                    run.italic = True

    def _add_table(self, block: dict[str, Any]) -> None:
        """Add a table to the document."""
        if not self.doc:
            return

        headers = block.get("headers", [])
        rows = block.get("rows", [])

        if not headers:
            return

        # Create table with header + data rows
        table = self.doc.add_table(rows=1 + len(rows), cols=len(headers))
        table.style = "Light Grid Accent 1"

        # Add header row
        header_cells = table.rows[0].cells
        for i, header_text in enumerate(headers):
            header_cells[i].text = header_text
            # Make header bold
            for paragraph in header_cells[i].paragraphs:
                for run in paragraph.runs:
                    run.bold = True

        # Add data rows
        for row_idx, row_data in enumerate(rows):
            row_cells = table.rows[row_idx + 1].cells
            for col_idx, cell_text in enumerate(row_data):
                if col_idx < len(row_cells):
                    row_cells[col_idx].text = str(cell_text)
