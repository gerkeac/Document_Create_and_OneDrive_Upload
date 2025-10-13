"""Tests for PowerPoint presentation generation."""

import tempfile
from collections.abc import Generator
from pathlib import Path

import pytest
from mcp_document_server.generators.powerpoint_generator import PowerPointGenerator
from pptx import Presentation


@pytest.fixture
def generator() -> PowerPointGenerator:
    """Create a PowerPoint generator instance."""
    return PowerPointGenerator()


@pytest.fixture
def temp_output_path() -> Generator[Path, None, None]:
    """Create a temporary output path for testing."""
    with tempfile.NamedTemporaryFile(suffix=".pptx", delete=False) as tmp:
        yield Path(tmp.name)


def test_generator_initialization(generator: PowerPointGenerator) -> None:
    """Test PowerPoint generator initialization."""
    assert generator is not None
    assert hasattr(generator, "generate_presentation")
    assert hasattr(generator, "logger")


def test_generate_title_slide(generator: PowerPointGenerator, temp_output_path: Path) -> None:
    """Test generating a presentation with a title slide."""
    slides = [
        {
            "layout": "title",
            "title": "Test Presentation",
            "subtitle": "Created with MCP Server",
        }
    ]

    result_path = generator.generate_presentation(slides, temp_output_path)

    assert result_path.exists()
    assert result_path.suffix == ".pptx"

    # Verify presentation contents
    prs = Presentation(str(result_path))
    assert len(prs.slides) == 1

    slide = prs.slides[0]
    assert slide.shapes.title.text == "Test Presentation"
    assert slide.placeholders[1].text == "Created with MCP Server"

    # Cleanup
    result_path.unlink()


def test_generate_title_content_slide(
    generator: PowerPointGenerator, temp_output_path: Path
) -> None:
    """Test generating a slide with title and bulleted content."""
    slides = [
        {
            "layout": "title_content",
            "title": "Key Features",
            "content": [
                "Document generation",
                "OneDrive integration",
                "Multi-user support",
            ],
        }
    ]

    result_path = generator.generate_presentation(slides, temp_output_path)

    assert result_path.exists()

    # Verify presentation contents
    prs = Presentation(str(result_path))
    assert len(prs.slides) == 1

    slide = prs.slides[0]
    assert slide.shapes.title.text == "Key Features"

    # Verify content bullets
    text_frame = slide.placeholders[1].text_frame
    paragraphs = [p.text for p in text_frame.paragraphs if p.text]
    assert "Document generation" in paragraphs
    assert "OneDrive integration" in paragraphs
    assert "Multi-user support" in paragraphs

    # Cleanup
    result_path.unlink()


def test_generate_title_table_slide(generator: PowerPointGenerator, temp_output_path: Path) -> None:
    """Test generating a slide with a table."""
    slides = [
        {
            "layout": "title_table",
            "title": "Q4 Results",
            "table": {
                "headers": ["Month", "Revenue", "Growth"],
                "rows": [
                    ["October", "$100K", "10%"],
                    ["November", "$120K", "20%"],
                    ["December", "$150K", "25%"],
                ],
            },
        }
    ]

    result_path = generator.generate_presentation(slides, temp_output_path)

    assert result_path.exists()

    # Verify presentation contents
    prs = Presentation(str(result_path))
    assert len(prs.slides) == 1

    slide = prs.slides[0]
    assert slide.shapes.title.text == "Q4 Results"

    # Find table in slide shapes
    table = None
    for shape in slide.shapes:
        if shape.has_table:
            table = shape.table
            break

    assert table is not None
    assert table.rows[0].cells[0].text == "Month"
    assert table.rows[1].cells[1].text == "$100K"
    assert table.rows[3].cells[2].text == "25%"

    # Cleanup
    result_path.unlink()


def test_generate_title_two_columns_slide(
    generator: PowerPointGenerator, temp_output_path: Path
) -> None:
    """Test generating a slide with two columns."""
    slides = [
        {
            "layout": "title_two_columns",
            "title": "Comparison",
            "left_content": ["Pro 1", "Pro 2", "Pro 3"],
            "right_content": ["Con 1", "Con 2", "Con 3"],
        }
    ]

    result_path = generator.generate_presentation(slides, temp_output_path)

    assert result_path.exists()

    # Verify presentation contents
    prs = Presentation(str(result_path))
    assert len(prs.slides) == 1

    slide = prs.slides[0]
    # Title should be in first textbox
    textboxes = [shape for shape in slide.shapes if hasattr(shape, "text_frame")]
    assert len(textboxes) >= 3  # Title + 2 columns

    # Cleanup
    result_path.unlink()


def test_generate_blank_slide(generator: PowerPointGenerator, temp_output_path: Path) -> None:
    """Test generating a blank slide with custom content."""
    slides = [
        {
            "layout": "blank",
            "title": "Custom Slide",
            "content": ["Custom content line 1", "Custom content line 2"],
        }
    ]

    result_path = generator.generate_presentation(slides, temp_output_path)

    assert result_path.exists()

    # Verify presentation contents
    prs = Presentation(str(result_path))
    assert len(prs.slides) == 1

    # Cleanup
    result_path.unlink()


def test_generate_multi_slide_presentation(
    generator: PowerPointGenerator, temp_output_path: Path
) -> None:
    """Test generating a presentation with multiple slides."""
    from typing import Any

    slides: list[dict[str, Any]] = [
        {
            "layout": "title",
            "title": "Project Overview",
            "subtitle": "Q4 2025",
        },
        {
            "layout": "title_content",
            "title": "Objectives",
            "content": ["Goal 1", "Goal 2", "Goal 3"],
        },
        {
            "layout": "title_table",
            "title": "Timeline",
            "table": {
                "headers": ["Phase", "Duration", "Status"],
                "rows": [
                    ["Phase 1", "2 weeks", "Complete"],
                    ["Phase 2", "2 weeks", "In Progress"],
                ],
            },
        },
    ]

    result_path = generator.generate_presentation(slides, temp_output_path)

    assert result_path.exists()

    # Verify presentation contents
    prs = Presentation(str(result_path))
    assert len(prs.slides) == 3

    # Verify each slide
    assert prs.slides[0].shapes.title.text == "Project Overview"
    assert prs.slides[1].shapes.title.text == "Objectives"
    assert prs.slides[2].shapes.title.text == "Timeline"

    # Cleanup
    result_path.unlink()


def test_generate_with_theme(generator: PowerPointGenerator, temp_output_path: Path) -> None:
    """Test generating a presentation with a specific theme."""
    slides = [
        {
            "layout": "title",
            "title": "Themed Presentation",
            "subtitle": "Blue Theme",
        }
    ]

    result_path = generator.generate_presentation(slides, temp_output_path, theme="blue")

    assert result_path.exists()

    # Verify presentation was created
    prs = Presentation(str(result_path))
    assert len(prs.slides) == 1

    # Cleanup
    result_path.unlink()


def test_empty_slides_list(generator: PowerPointGenerator, temp_output_path: Path) -> None:
    """Test handling of empty slides list."""
    from typing import Any

    slides: list[dict[str, Any]] = []

    # Empty slides should create a presentation with no slides
    result_path = generator.generate_presentation(slides, temp_output_path)

    assert result_path.exists()

    # Verify presentation has no slides
    prs = Presentation(str(result_path))
    assert len(prs.slides) == 0

    # Cleanup
    result_path.unlink()


def test_invalid_layout_fallback(generator: PowerPointGenerator, temp_output_path: Path) -> None:
    """Test that invalid layout types fall back to title_content."""
    slides = [
        {
            "layout": "invalid_layout_type",
            "title": "Fallback Test",
            "content": ["This should use title_content layout"],
        }
    ]

    result_path = generator.generate_presentation(slides, temp_output_path)

    assert result_path.exists()

    # Verify presentation was created with fallback
    prs = Presentation(str(result_path))
    assert len(prs.slides) == 1

    # Cleanup
    result_path.unlink()


def test_nested_bullet_content(generator: PowerPointGenerator, temp_output_path: Path) -> None:
    """Test generating slides with nested bullet points."""
    slides = [
        {
            "layout": "title_content",
            "title": "Nested Bullets",
            "content": [
                {"text": "Main point 1", "sub_items": ["Sub 1.1", "Sub 1.2"]},
                {"text": "Main point 2", "sub_items": ["Sub 2.1"]},
            ],
        }
    ]

    result_path = generator.generate_presentation(slides, temp_output_path)

    assert result_path.exists()

    # Verify presentation contents
    prs = Presentation(str(result_path))
    assert len(prs.slides) == 1

    # Cleanup
    result_path.unlink()


def test_string_content_conversion(generator: PowerPointGenerator, temp_output_path: Path) -> None:
    """Test that string content is converted to list."""
    slides = [
        {
            "layout": "title_content",
            "title": "Single String",
            "content": "This is a single string content",
        }
    ]

    result_path = generator.generate_presentation(slides, temp_output_path)

    assert result_path.exists()

    # Verify presentation contents
    prs = Presentation(str(result_path))
    assert len(prs.slides) == 1

    # Cleanup
    result_path.unlink()


def test_output_directory_creation(generator: PowerPointGenerator) -> None:
    """Test that output directory is created if it doesn't exist."""
    with tempfile.TemporaryDirectory() as tmpdir:
        output_path = Path(tmpdir) / "subdir" / "presentation.pptx"

        slides = [
            {
                "layout": "title",
                "title": "Test",
                "subtitle": "Directory Creation",
            }
        ]

        result_path = generator.generate_presentation(slides, output_path)

        assert result_path.exists()
        assert result_path.parent.exists()

        # Cleanup
        result_path.unlink()
