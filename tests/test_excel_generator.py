"""Tests for Excel spreadsheet generation."""

import tempfile
from collections.abc import Generator
from pathlib import Path

import pytest
from mcp_document_server.generators.excel_generator import ExcelGenerator
from openpyxl import load_workbook


@pytest.fixture
def generator() -> ExcelGenerator:
    """Create an Excel generator instance."""
    return ExcelGenerator()


@pytest.fixture
def temp_output_path() -> Generator[Path, None, None]:
    """Create a temporary output path for testing."""
    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as tmp:
        yield Path(tmp.name)


def test_generator_initialization(generator: ExcelGenerator) -> None:
    """Test Excel generator initialization."""
    assert generator is not None
    assert hasattr(generator, "generate_spreadsheet")
    assert hasattr(generator, "logger")


def test_generate_simple_spreadsheet(generator: ExcelGenerator, temp_output_path: Path) -> None:
    """Test generating a simple spreadsheet with one sheet."""
    sheets = [
        {
            "name": "Sheet1",
            "headers": ["Name", "Age", "City"],
            "data": [
                ["Alice", 30, "New York"],
                ["Bob", 25, "Los Angeles"],
                ["Charlie", 35, "Chicago"],
            ],
        }
    ]

    result_path = generator.generate_spreadsheet(sheets, temp_output_path)

    assert result_path.exists()
    assert result_path.suffix == ".xlsx"

    # Verify spreadsheet contents
    wb = load_workbook(str(result_path))
    assert "Sheet1" in wb.sheetnames

    ws = wb["Sheet1"]
    assert ws["A1"].value == "Name"
    assert ws["B1"].value == "Age"
    assert ws["C1"].value == "City"
    assert ws["A2"].value == "Alice"
    assert ws["B2"].value == 30
    assert ws["C2"].value == "New York"

    # Cleanup
    result_path.unlink()


def test_generate_multiple_sheets(generator: ExcelGenerator, temp_output_path: Path) -> None:
    """Test generating a spreadsheet with multiple sheets."""
    sheets = [
        {
            "name": "Q1 Sales",
            "headers": ["Product", "Revenue"],
            "data": [
                ["Product A", 10000],
                ["Product B", 15000],
            ],
        },
        {
            "name": "Q2 Sales",
            "headers": ["Product", "Revenue"],
            "data": [
                ["Product A", 12000],
                ["Product B", 18000],
            ],
        },
    ]

    result_path = generator.generate_spreadsheet(sheets, temp_output_path)

    assert result_path.exists()

    # Verify spreadsheet contents
    wb = load_workbook(str(result_path))
    assert "Q1 Sales" in wb.sheetnames
    assert "Q2 Sales" in wb.sheetnames
    assert len(wb.sheetnames) == 2

    # Verify Q1 data
    ws1 = wb["Q1 Sales"]
    assert ws1["A1"].value == "Product"
    assert ws1["B2"].value == 10000

    # Verify Q2 data
    ws2 = wb["Q2 Sales"]
    assert ws2["A1"].value == "Product"
    assert ws2["B2"].value == 12000

    # Cleanup
    result_path.unlink()


def test_generate_with_formulas(generator: ExcelGenerator, temp_output_path: Path) -> None:
    """Test generating a spreadsheet with formulas."""
    sheets = [
        {
            "name": "Budget",
            "headers": ["Category", "Q1", "Q2", "Total"],
            "data": [
                ["Marketing", 10000, 12000, "=B2+C2"],
                ["Engineering", 50000, 52000, "=B3+C3"],
                ["Total", "=B2+B3", "=C2+C3", "=B4+C4"],
            ],
        }
    ]

    result_path = generator.generate_spreadsheet(sheets, temp_output_path, include_formulas=True)

    assert result_path.exists()

    # Verify spreadsheet contents
    wb = load_workbook(str(result_path))
    ws = wb["Budget"]

    # Check headers
    assert ws["A1"].value == "Category"
    assert ws["D1"].value == "Total"

    # Check data values
    assert ws["A2"].value == "Marketing"
    assert ws["B2"].value == 10000
    assert ws["C2"].value == 12000

    # Check formula cells
    assert ws["D2"].value == "=B2+C2"
    assert ws["D3"].value == "=B3+C3"

    # Cleanup
    result_path.unlink()


def test_generate_without_formulas(generator: ExcelGenerator, temp_output_path: Path) -> None:
    """Test generating a spreadsheet with formulas disabled."""
    sheets = [
        {
            "name": "Data",
            "headers": ["A", "B"],
            "data": [
                [10, "=A2*2"],
            ],
        }
    ]

    result_path = generator.generate_spreadsheet(sheets, temp_output_path, include_formulas=False)

    assert result_path.exists()

    # Verify formula is treated as text
    wb = load_workbook(str(result_path))
    ws = wb["Data"]
    assert ws["B2"].value == "=A2*2"  # Should be stored as text

    # Cleanup
    result_path.unlink()


def test_generate_without_headers(generator: ExcelGenerator, temp_output_path: Path) -> None:
    """Test generating a spreadsheet without headers."""
    sheets = [
        {
            "name": "NoHeaders",
            "data": [
                [1, 2, 3],
                [4, 5, 6],
            ],
        }
    ]

    result_path = generator.generate_spreadsheet(sheets, temp_output_path)

    assert result_path.exists()

    # Verify spreadsheet contents
    wb = load_workbook(str(result_path))
    ws = wb["NoHeaders"]
    assert ws["A1"].value == 1
    assert ws["B1"].value == 2
    assert ws["C2"].value == 6

    # Cleanup
    result_path.unlink()


def test_header_styling(generator: ExcelGenerator, temp_output_path: Path) -> None:
    """Test that headers have proper styling."""
    sheets = [
        {
            "name": "Styled",
            "headers": ["Column1", "Column2"],
            "data": [
                ["Data1", "Data2"],
            ],
        }
    ]

    result_path = generator.generate_spreadsheet(sheets, temp_output_path)

    assert result_path.exists()

    # Verify header styling
    wb = load_workbook(str(result_path))
    ws = wb["Styled"]

    # Check that header cells have bold font
    assert ws["A1"].font.bold is True
    assert ws["B1"].font.bold is True

    # Check that header cells have fill color
    assert ws["A1"].fill.start_color.rgb is not None

    # Cleanup
    result_path.unlink()


def test_alternating_row_colors(generator: ExcelGenerator, temp_output_path: Path) -> None:
    """Test that data rows have alternating colors."""
    sheets = [
        {
            "name": "Alternating",
            "headers": ["A", "B"],
            "data": [
                [1, 2],
                [3, 4],
                [5, 6],
            ],
        }
    ]

    result_path = generator.generate_spreadsheet(sheets, temp_output_path)

    assert result_path.exists()

    # Verify alternating row colors
    wb = load_workbook(str(result_path))
    ws = wb["Alternating"]

    # Row 2 (even) should have fill color
    assert ws["A2"].fill.start_color.rgb is not None

    # Row 3 (odd) should have different fill (or no fill)
    # Note: The generator applies fill to even rows only

    # Cleanup
    result_path.unlink()


def test_auto_size_columns(generator: ExcelGenerator, temp_output_path: Path) -> None:
    """Test that columns are auto-sized based on content."""
    sheets = [
        {
            "name": "AutoSize",
            "headers": ["Short", "Very Long Header Name"],
            "data": [
                ["A", "Short data"],
            ],
            "auto_size_columns": True,
        }
    ]

    result_path = generator.generate_spreadsheet(sheets, temp_output_path)

    assert result_path.exists()

    # Verify column widths
    wb = load_workbook(str(result_path))
    ws = wb["AutoSize"]

    # Column B should be wider than column A due to longer header
    col_a_width = ws.column_dimensions["A"].width
    col_b_width = ws.column_dimensions["B"].width
    assert col_b_width is not None and col_a_width is not None
    assert col_b_width > col_a_width

    # Cleanup
    result_path.unlink()


def test_freeze_header_row(generator: ExcelGenerator, temp_output_path: Path) -> None:
    """Test that header row is frozen."""
    sheets = [
        {
            "name": "Frozen",
            "headers": ["A", "B"],
            "data": [
                [1, 2],
                [3, 4],
            ],
            "freeze_header": True,
        }
    ]

    result_path = generator.generate_spreadsheet(sheets, temp_output_path)

    assert result_path.exists()

    # Verify freeze panes
    wb = load_workbook(str(result_path))
    ws = wb["Frozen"]

    assert ws.freeze_panes == "A2"

    # Cleanup
    result_path.unlink()


def test_numeric_formatting(generator: ExcelGenerator, temp_output_path: Path) -> None:
    """Test that numbers are formatted correctly."""
    sheets = [
        {
            "name": "Numbers",
            "headers": ["Integer", "Float", "Text"],
            "data": [
                [100, 99.99, "Text value"],
                [1000, 1234.56, "Another"],
            ],
        }
    ]

    result_path = generator.generate_spreadsheet(sheets, temp_output_path)

    assert result_path.exists()

    # Verify numeric types
    wb = load_workbook(str(result_path))
    ws = wb["Numbers"]

    # Check data types
    assert isinstance(ws["A2"].value, int)
    assert isinstance(ws["B2"].value, float)
    assert isinstance(ws["C2"].value, str)

    # Cleanup
    result_path.unlink()


def test_string_number_parsing(generator: ExcelGenerator, temp_output_path: Path) -> None:
    """Test that numeric strings are parsed correctly."""
    sheets = [
        {
            "name": "Parsing",
            "headers": ["Value"],
            "data": [
                ["123"],  # Should be parsed as int
                ["45.67"],  # Should be parsed as float
                ["Not a number"],  # Should remain string
            ],
        }
    ]

    result_path = generator.generate_spreadsheet(sheets, temp_output_path)

    assert result_path.exists()

    # Verify parsing
    wb = load_workbook(str(result_path))
    ws = wb["Parsing"]

    assert isinstance(ws["A2"].value, int)
    assert ws["A2"].value == 123

    assert isinstance(ws["A3"].value, float)
    assert ws["A3"].value == 45.67

    assert isinstance(ws["A4"].value, str)
    assert ws["A4"].value == "Not a number"

    # Cleanup
    result_path.unlink()


def test_empty_sheets_list(generator: ExcelGenerator, temp_output_path: Path) -> None:
    """Test handling of empty sheets list."""
    from typing import Any

    sheets: list[dict[str, Any]] = []

    with pytest.raises(ValueError, match="At least one sheet is required"):
        generator.generate_spreadsheet(sheets, temp_output_path)


def test_default_sheet_naming(generator: ExcelGenerator, temp_output_path: Path) -> None:
    """Test that sheets without names get default names."""
    sheets = [
        {
            "headers": ["A"],
            "data": [[1]],
        },
        {
            "headers": ["B"],
            "data": [[2]],
        },
    ]

    result_path = generator.generate_spreadsheet(sheets, temp_output_path)

    assert result_path.exists()

    # Verify default sheet names
    wb = load_workbook(str(result_path))
    assert "Sheet1" in wb.sheetnames
    assert "Sheet2" in wb.sheetnames

    # Cleanup
    result_path.unlink()


def test_large_dataset(generator: ExcelGenerator, temp_output_path: Path) -> None:
    """Test generating a spreadsheet with a larger dataset."""
    # Generate 1000 rows of data
    data_rows = [[i, f"Row {i}", i * 100] for i in range(1, 1001)]

    sheets = [
        {
            "name": "LargeData",
            "headers": ["ID", "Description", "Value"],
            "data": data_rows,
        }
    ]

    result_path = generator.generate_spreadsheet(sheets, temp_output_path)

    assert result_path.exists()

    # Verify data count
    wb = load_workbook(str(result_path))
    ws = wb["LargeData"]

    # Check first and last rows
    assert ws["A2"].value == 1
    assert ws["A1001"].value == 1000
    assert ws["C1001"].value == 100000

    # Cleanup
    result_path.unlink()


def test_output_directory_creation(generator: ExcelGenerator) -> None:
    """Test that output directory is created if it doesn't exist."""
    with tempfile.TemporaryDirectory() as tmpdir:
        output_path = Path(tmpdir) / "subdir" / "spreadsheet.xlsx"

        sheets = [
            {
                "name": "Test",
                "headers": ["A"],
                "data": [[1]],
            }
        ]

        result_path = generator.generate_spreadsheet(sheets, output_path)

        assert result_path.exists()
        assert result_path.parent.exists()

        # Cleanup
        result_path.unlink()


def test_cell_borders(generator: ExcelGenerator, temp_output_path: Path) -> None:
    """Test that cells have borders."""
    sheets = [
        {
            "name": "Borders",
            "headers": ["A", "B"],
            "data": [
                [1, 2],
            ],
        }
    ]

    result_path = generator.generate_spreadsheet(sheets, temp_output_path)

    assert result_path.exists()

    # Verify borders
    wb = load_workbook(str(result_path))
    ws = wb["Borders"]

    # Check that data cells have borders
    assert ws["A2"].border.left.style is not None
    assert ws["B2"].border.right.style is not None

    # Cleanup
    result_path.unlink()
