"""Tests for Word document generation."""

import tempfile
from pathlib import Path

import pytest
from docx import Document
from mcp_document_server.generators.word_generator import WordGenerator


@pytest.fixture
def temp_output_path() -> Path:  # type: ignore[misc]
    """Create a temporary output path for testing."""
    temp_dir = Path(tempfile.gettempdir()) / "test_mcp_documents"
    temp_dir.mkdir(exist_ok=True)
    output_path = temp_dir / "test_document.docx"

    yield output_path

    # Cleanup
    if output_path.exists():
        output_path.unlink()


def test_word_generator_initialization() -> None:
    """Test that WordGenerator can be initialized."""
    generator = WordGenerator()
    assert generator is not None
    assert generator.doc is None


def test_generate_from_markdown_simple(temp_output_path: Path) -> None:
    """Test generating Word document from simple markdown."""
    generator = WordGenerator()
    markdown = "# Test Document\n\nThis is a test paragraph."

    result = generator.generate_from_markdown(
        content=markdown,
        output_path=temp_output_path,
    )

    assert result == temp_output_path
    assert temp_output_path.exists()

    # Verify document content
    doc = Document(str(temp_output_path))
    assert len(doc.paragraphs) >= 2


def test_generate_from_markdown_with_metadata(temp_output_path: Path) -> None:
    """Test generating Word document with metadata."""
    generator = WordGenerator()
    markdown = "# Test Document\n\nContent here."

    result = generator.generate_from_markdown(
        content=markdown,
        output_path=temp_output_path,
        title="Test Title",
        author="Test Author",
    )

    assert result.exists()

    # Verify metadata
    doc = Document(str(temp_output_path))
    assert doc.core_properties.title == "Test Title"
    assert doc.core_properties.author == "Test Author"


def test_generate_from_markdown_headings(temp_output_path: Path) -> None:
    """Test generating Word document with multiple heading levels."""
    generator = WordGenerator()
    markdown = """# Heading 1
## Heading 2
### Heading 3"""

    generator.generate_from_markdown(
        content=markdown,
        output_path=temp_output_path,
    )

    doc = Document(str(temp_output_path))
    paragraphs = doc.paragraphs

    # Check that headings are created
    assert len(paragraphs) >= 3
    assert "Heading 1" in paragraphs[0].text
    assert "Heading 2" in paragraphs[1].text
    assert "Heading 3" in paragraphs[2].text


def test_generate_from_markdown_lists(temp_output_path: Path) -> None:
    """Test generating Word document with lists."""
    generator = WordGenerator()
    markdown = """- Item 1
- Item 2
- Item 3

1. First
2. Second
3. Third"""

    generator.generate_from_markdown(
        content=markdown,
        output_path=temp_output_path,
    )

    doc = Document(str(temp_output_path))
    assert temp_output_path.exists()

    # Check that list items exist
    paragraphs = doc.paragraphs
    list_items = [p for p in paragraphs if "List" in str(p.style.name)]
    assert len(list_items) >= 6  # 3 bullet + 3 numbered


def test_generate_from_markdown_table(temp_output_path: Path) -> None:
    """Test generating Word document with table."""
    generator = WordGenerator()
    markdown = """| Name | Age | City |
|------|-----|------|
| Alice | 30 | NYC |
| Bob | 25 | LA |"""

    generator.generate_from_markdown(
        content=markdown,
        output_path=temp_output_path,
    )

    doc = Document(str(temp_output_path))
    assert len(doc.tables) == 1

    table = doc.tables[0]
    assert len(table.rows) == 3  # Header + 2 data rows
    assert len(table.columns) == 3


def test_generate_from_json_simple(temp_output_path: Path) -> None:
    """Test generating Word document from JSON."""
    generator = WordGenerator()
    content = {
        "title": "JSON Document",
        "author": "Test Author",
        "blocks": [
            {"type": "heading", "level": 1, "text": "Main Title"},
            {"type": "paragraph", "text": "This is a paragraph."},
        ],
    }

    result = generator.generate_from_json(
        content=content,
        output_path=temp_output_path,
    )

    assert result.exists()

    doc = Document(str(temp_output_path))
    assert doc.core_properties.title == "JSON Document"
    assert doc.core_properties.author == "Test Author"


def test_generate_from_json_with_list(temp_output_path: Path) -> None:
    """Test generating Word document from JSON with lists."""
    generator = WordGenerator()
    content = {
        "blocks": [
            {
                "type": "list",
                "list_type": "bullet",
                "items": ["Item A", "Item B", "Item C"],
            },
        ],
    }

    generator.generate_from_json(
        content=content,
        output_path=temp_output_path,
    )

    doc = Document(str(temp_output_path))
    list_paragraphs = [p for p in doc.paragraphs if "List" in str(p.style.name)]
    assert len(list_paragraphs) == 3


def test_generate_from_json_with_table(temp_output_path: Path) -> None:
    """Test generating Word document from JSON with table."""
    generator = WordGenerator()
    content = {
        "blocks": [
            {
                "type": "table",
                "headers": ["Col1", "Col2"],
                "rows": [["A", "B"], ["C", "D"]],
            },
        ],
    }

    generator.generate_from_json(
        content=content,
        output_path=temp_output_path,
    )

    doc = Document(str(temp_output_path))
    assert len(doc.tables) == 1
    table = doc.tables[0]
    assert len(table.rows) == 3  # Header + 2 rows


def test_generate_creates_parent_directory(temp_output_path: Path) -> None:
    """Test that generator creates parent directories if needed."""
    generator = WordGenerator()
    nested_path = temp_output_path.parent / "nested" / "deep" / "document.docx"

    markdown = "# Test"
    result = generator.generate_from_markdown(
        content=markdown,
        output_path=nested_path,
    )

    assert result.exists()
    assert result.parent.exists()

    # Cleanup
    nested_path.unlink()
    nested_path.parent.rmdir()
    nested_path.parent.parent.rmdir()


def test_generate_from_markdown_inline_formatting(temp_output_path: Path) -> None:
    """Test generating Word document with inline formatting."""
    generator = WordGenerator()
    markdown = "This is **bold** and *italic* text."

    generator.generate_from_markdown(
        content=markdown,
        output_path=temp_output_path,
    )

    doc = Document(str(temp_output_path))
    paragraph = doc.paragraphs[0]

    # Check that runs exist with formatting
    assert len(paragraph.runs) > 1

    # Find bold run
    bold_runs = [r for r in paragraph.runs if r.bold]
    assert len(bold_runs) >= 1

    # Find italic run
    italic_runs = [r for r in paragraph.runs if r.italic]
    assert len(italic_runs) >= 1


def test_generate_from_markdown_complex(temp_output_path: Path) -> None:
    """Test generating complex Word document."""
    generator = WordGenerator()
    markdown = """# Annual Report 2025

## Executive Summary

This report shows **strong growth** in all areas.

## Key Metrics

- Revenue: **$1.2M**
- Growth: *25%*
- Users: ***100,000***

## Data Table

| Quarter | Revenue | Growth |
|---------|---------|--------|
| Q1      | $250K   | 20%    |
| Q2      | $300K   | 25%    |
| Q3      | $325K   | 30%    |
| Q4      | $325K   | 25%    |

## Next Steps

1. Expand team
2. Launch new product
3. Enter new markets"""

    result = generator.generate_from_markdown(
        content=markdown,
        output_path=temp_output_path,
        title="Annual Report",
        author="Test Company",
    )

    assert result.exists()

    doc = Document(str(temp_output_path))
    assert doc.core_properties.title == "Annual Report"
    assert len(doc.paragraphs) > 10
    assert len(doc.tables) == 1
